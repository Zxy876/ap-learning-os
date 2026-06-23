#!/usr/bin/env python3
import argparse
import datetime as dt
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

from import_compile_snapshot import connect


BASE = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT_DIR = BASE / "data" / "online_platform" / "writebacks"
WRITEBACK_HEADERS = {
    "review_status": "APLOS Review Status",
    "failure_type": "APLOS Failure Type",
    "review_message": "APLOS Review Message",
    "reviewed_at": "APLOS Reviewed At",
    "task_date": "APLOS Task Date",
    "task_title": "APLOS Task Title",
}


def load_json(value, fallback):
    try:
        return json.loads(value) if value else fallback
    except json.JSONDecodeError:
        return fallback


def ensure_header(ws, header, header_row=3):
    for cell in ws[header_row]:
        if str(cell.value or "").strip() == header:
            return cell.column
    col = ws.max_column + 1
    ws.cell(row=header_row, column=col).value = header
    return col


def workbook_output_path(source_path, output_dir):
    source = Path(source_path)
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    return Path(output_dir) / f"{source.stem}.aplos-writeback-{stamp}{source.suffix}"


def pending_writebacks(conn):
    return conn.execute(
        """
        SELECT sw.*, pr.source_workbook, pr.source_sheet, pr.source_row_number
        FROM spreadsheet_writebacks sw
        LEFT JOIN plan_rows pr ON pr.id = sw.plan_row_id
        WHERE sw.status = 'pending'
        ORDER BY sw.created_at ASC
        """
    ).fetchall()


def writeback_groups(rows):
    groups = {}
    for row in rows:
        workbook = row["source_workbook"]
        if not workbook:
            continue
        groups.setdefault(workbook, []).append(row)
    return groups


def apply_rows_to_workbook(conn, source_path, rows, output_dir, in_place=False):
    source = Path(source_path)
    if not source.exists():
        for row in rows:
            conn.execute(
                "UPDATE spreadsheet_writebacks SET status = ?, payload_json = ? WHERE id = ?",
                ("failed", json.dumps({"error": f"source workbook missing: {source}"}, ensure_ascii=False), row["id"]),
            )
        return {"source": str(source), "status": "failed", "error": "source workbook missing", "rows": len(rows)}

    output = source if in_place else workbook_output_path(source, output_dir)
    if not in_place:
        output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, output)

    wb = load_workbook(output)
    written = 0
    for row in rows:
        sheet = row["target_sheet"] or row["source_sheet"]
        target_row = row["target_row"] or row["source_row_number"]
        if not sheet or sheet not in wb.sheetnames or not target_row:
            conn.execute(
                "UPDATE spreadsheet_writebacks SET status = ?, payload_json = ? WHERE id = ?",
                ("failed", json.dumps({"error": "missing sheet or target row"}, ensure_ascii=False), row["id"]),
            )
            continue
        ws = wb[sheet]
        payload = load_json(row["payload_json"], {})
        values = {
            "review_status": payload.get("final_state", ""),
            "failure_type": payload.get("failure_type", ""),
            "review_message": payload.get("message", ""),
            "reviewed_at": dt.datetime.now().isoformat(timespec="seconds"),
            "task_date": payload.get("scheduled_date", ""),
            "task_title": payload.get("title", ""),
        }
        for key, header in WRITEBACK_HEADERS.items():
            col = ensure_header(ws, header)
            ws.cell(row=int(target_row), column=col).value = values.get(key, "")
        conn.execute(
            "UPDATE spreadsheet_writebacks SET status = ?, payload_json = ? WHERE id = ?",
            (
                "written",
                json.dumps({**payload, "writeback_workbook": str(output), "written_at": values["reviewed_at"]}, ensure_ascii=False),
                row["id"],
            ),
        )
        written += 1
    wb.save(output)
    return {"source": str(source), "output": str(output), "status": "written", "rows": written}


def process_writebacks(conn, output_dir=DEFAULT_OUTPUT_DIR, in_place=False):
    output_dir = output_dir or DEFAULT_OUTPUT_DIR
    rows = pending_writebacks(conn)
    result = {"pending": len(rows), "workbooks": []}
    for source_path, group in writeback_groups(rows).items():
        result["workbooks"].append(apply_rows_to_workbook(conn, source_path, group, output_dir, in_place))
    conn.commit()
    result["written"] = sum(item.get("rows", 0) for item in result["workbooks"] if item.get("status") == "written")
    result["failed"] = conn.execute("SELECT COUNT(*) FROM spreadsheet_writebacks WHERE status = 'failed'").fetchone()[0]
    return result


def main():
    parser = argparse.ArgumentParser(description="Apply pending AP Learning OS review writebacks to Excel workbooks.")
    parser.add_argument("--db", default=str(BASE / "data" / "online_platform" / "aplos_dev.sqlite3"))
    parser.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR))
    parser.add_argument("--in-place", action="store_true", help="write directly into source workbooks instead of a derived copy")
    args = parser.parse_args()
    conn = connect(args.db)
    print(json.dumps(process_writebacks(conn, args.output_dir, args.in_place), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
