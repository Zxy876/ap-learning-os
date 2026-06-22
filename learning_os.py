#!/usr/bin/env python3
import argparse
import datetime as dt
import hashlib
import json
import logging
import os
import re
import shutil
import contextlib
import subprocess
import sys
import time
import uuid
from pathlib import Path

from openpyxl import load_workbook


BASE = Path(__file__).resolve().parent
CONFIG_PATH = BASE / "config.json"
STATE_PATH = BASE / "data" / "state.json"
BLACKBOARD_PATH = BASE / "blackboard" / "today.md"
TASK_MATERIALS_DIR = BASE / "task_materials"
RESOURCE_RESOLUTION_CACHE = {}

logging.getLogger("pypdf").setLevel(logging.ERROR)


TASK_STATES = {
    "Planned",
    "Running",
    "Candidate Complete",
    "Needs Review",
    "Completed",
    "Partially Completed",
    "Failed",
    "Blocked",
}


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def load_config():
    return read_json(CONFIG_PATH)


def load_state():
    if not STATE_PATH.exists():
        return {"tasks": {}, "sessions": {}, "adaptive_time": {}}
    return read_json(STATE_PATH)


def save_state(state):
    write_json(STATE_PATH, state)


def parse_date_arg(value):
    return dt.datetime.strptime(value, "%Y-%m-%d").date()


def today_date():
    return dt.date.today()


def slug(value):
    value = re.sub(r"[^A-Za-z0-9一-龥]+", "-", str(value)).strip("-")
    return value[:80] or "task"


def row_dicts(ws, header_row):
    headers = [cell.value for cell in ws[header_row]]
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(v not in (None, "") for v in row):
            continue
        yield {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}


def normalize_plan_mmdd(raw_date, year):
    if raw_date in (None, ""):
        return None
    if isinstance(raw_date, dt.datetime):
        return raw_date.date()
    if isinstance(raw_date, dt.date):
        return raw_date
    text = str(raw_date).strip()
    m = re.match(r"^(\d{1,2})/(\d{1,2})$", text)
    if not m:
        return None
    return dt.date(year, int(m.group(1)), int(m.group(2)))


def parse_day_range(raw):
    text = str(raw or "")
    nums = [int(n) for n in re.findall(r"\d+", text)]
    if not nums:
        return None
    if len(nums) == 1:
        return nums[0], nums[0]
    return nums[0], nums[1]


def adaptive_minutes(state, task_key, default_min, config):
    return int(state.get("adaptive_time", {}).get(task_key, default_min))


def find_existing_path(config, candidate):
    if not candidate or str(candidate).strip() in {"—", "-"}:
        return None
    text = str(candidate).strip()
    workspace = Path(config["workspace_root"])
    if Path(text).exists():
        return str(Path(text))
    basename = Path(text.replace("\\", "/")).name
    if not basename:
        return None
    matches = list(workspace.rglob(basename))
    if matches:
        return str(matches[0])
    return None


def abs_workspace_path(config, relative_path):
    if not relative_path:
        return None
    path = Path(relative_path)
    if not path.is_absolute():
        path = Path(config["workspace_root"]) / path
    return str(path) if path.exists() else None


def local_fallback_resources(config, course, unit, include_practice=False):
    fallbacks = config.get("local_resource_fallbacks", {}).get(course, {})
    resources = []
    for label in ["workbook", "dashboard", "summary"]:
        target = abs_workspace_path(config, fallbacks.get(label))
        if target:
            resources.append({"label": f"local_{label}", "target": target})
    unit_target = abs_workspace_path(config, fallbacks.get("unit_pages", {}).get(str(unit)))
    if unit_target:
        resources.append({"label": "local_unit_page", "target": unit_target})
    if include_practice:
        practice_target = abs_workspace_path(config, fallbacks.get("practice_pdfs", {}).get(str(unit)))
        if practice_target:
            resources.append({"label": "local_practice_pdf", "target": practice_target})
    return resources


def normalize_resource_key(value):
    text = str(value or "").lower().replace("…", "").replace("...", "")
    return re.sub(r"[^a-z0-9一-龥]+", "", text)


def resource_search_roots(config):
    roots = [
        Path(config["workspace_root"]),
        Path("/Users/zxydediannao/Downloads"),
        Path("/Users/zxydediannao/Library/Mobile Documents/com~apple~CloudDocs"),
        Path("/Users/zxydediannao/Library/Mobile Documents/iCloud~QReader~MarginStudy~easy/Documents"),
    ]
    clean = []
    seen = set()
    for root in roots:
        if root.exists() and root not in seen:
            seen.add(root)
            clean.append(root)
    return clean


def is_generated_material(path):
    try:
        Path(path).resolve().relative_to(TASK_MATERIALS_DIR.resolve())
        return True
    except ValueError:
        return False


def resource_aliases(config):
    canonical = config.get("canonical_materials", {})
    aliases = {
        "AP_Calculus_BC": {
            "James Stewart - Calculus Early Transcendentals 8th.pdf":
                canonical.get("AP_Calculus_BC", {}).get("stewart_textbook"),
        },
        "AP_CSA": {
            "JAVA Illuminated.pdf":
                canonical.get("AP_CSA", {}).get("java_textbook"),
        }
    }
    for course, course_aliases in config.get("resource_index_aliases", {}).items():
        aliases.setdefault(course, {}).update(course_aliases)
    return aliases


def is_url(target):
    return str(target or "").startswith(("http://", "https://"))


def target_exists(target):
    return bool(target) and (is_url(target) or Path(target).exists())


def resource_index_entries(config, course):
    plan = Path(config["plans"][course])
    sheet = "资源索引"
    wb = load_workbook(plan, data_only=True, read_only=True)
    ws = wb[sheet]
    headers = [cell.value for cell in ws[3]]
    entries = []
    path_col = "完整路径" if "完整路径" in headers else "路径"
    for row_number, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        if not any(v not in (None, "") for v in row):
            continue
        data = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        filename = str(data.get("文件名") or "").strip()
        if not filename:
            continue
        entries.append({
            "course": course,
            "row": row_number,
            "category": str(data.get("类别") or "").strip(),
            "filename": filename,
            "path_hint": str(data.get(path_col) or "").strip(),
            "unit": str(data.get("对应单元") or "").strip(),
            "note": str(data.get("说明") or "").strip(),
        })
    return entries


def unit_matches(resource_unit, task_unit):
    text = str(resource_unit or "").strip()
    task = str(task_unit or "").strip()
    if not text or text in {"全部", "Unit 1-10"}:
        return True
    if not task:
        return False
    if text == task:
        return True
    nums = [int(n) for n in re.findall(r"\d+", text)]
    task_nums = [int(n) for n in re.findall(r"\d+", task)]
    if not nums or not task_nums:
        return False
    task_num = task_nums[0]
    if "-" in text and len(nums) >= 2:
        return nums[0] <= task_num <= nums[1]
    return task_num in nums


def direct_resource_candidate(root, path_hint, filename):
    rel_parts = []
    if path_hint:
        rel_parts.extend(p for p in str(path_hint).replace("\\", "/").split("/") if p)
    clean_filename = str(filename).replace("\\", "/").strip("/")
    if clean_filename:
        rel_parts.append(Path(clean_filename).name)
    if not rel_parts:
        return None
    return root.joinpath(*rel_parts)


def file_tokens(filename):
    stem = Path(str(filename).replace("\\", "/")).stem.replace("...", " ")
    return [
        normalize_resource_key(token)
        for token in re.split(r"[^A-Za-z0-9一-龥]+", stem)
        if len(normalize_resource_key(token)) >= 3
    ]


def strict_resource_match(entry, candidate):
    if is_generated_material(candidate):
        return 0
    filename = entry["filename"].replace("\\", "/").rstrip("/")
    wants_dir = entry["filename"].endswith("\\") or entry["filename"].endswith("/")
    if wants_dir and not candidate.is_dir():
        return 0
    if not wants_dir and not candidate.is_file():
        return 0
    expected_suffix = Path(filename).suffix.lower()
    if expected_suffix and candidate.suffix.lower() != expected_suffix:
        return 0
    path_text = str(candidate).lower()
    if entry["course"] == "AP_Calculus_BC" and entry["category"] in {"真题", "模考"}:
        if "calculus" not in path_text and "calc" not in path_text:
            return 0
    if entry["course"] == "AP_CSA" and entry["category"] in {"真题", "模考"}:
        if "csa" not in path_text and "computer science" not in path_text:
            return 0
    expected = normalize_resource_key(Path(filename).name)
    actual = normalize_resource_key(candidate.name)
    if expected and actual == expected:
        return 1000
    if wants_dir and expected and expected in actual:
        return 800
    if "..." not in entry["filename"] and "…" not in entry["filename"]:
        return 0
    tokens = file_tokens(entry["filename"])
    if not tokens:
        return 0
    hits = sum(1 for token in tokens if token in actual)
    if hits == len(tokens):
        return 700 + hits
    return 0


def resolve_index_resource(config, entry):
    cache_key = (entry["course"], entry["filename"], entry.get("path_hint", ""))
    if cache_key in RESOURCE_RESOLUTION_CACHE:
        return RESOURCE_RESOLUTION_CACHE[cache_key]
    alias = resource_aliases(config).get(entry["course"], {}).get(entry["filename"])
    if alias and Path(alias).exists():
        RESOURCE_RESOLUTION_CACHE[cache_key] = str(Path(alias))
        return RESOURCE_RESOLUTION_CACHE[cache_key]
    for root in resource_search_roots(config):
        direct = direct_resource_candidate(root, entry.get("path_hint"), entry["filename"])
        if direct and direct.exists() and not is_generated_material(direct):
            RESOURCE_RESOLUTION_CACHE[cache_key] = str(direct)
            return RESOURCE_RESOLUTION_CACHE[cache_key]
    best = None
    best_score = 0
    for root in resource_search_roots(config):
        for candidate in root.rglob("*"):
            score = strict_resource_match(entry, candidate)
            if score > best_score:
                best = candidate
                best_score = score
    RESOURCE_RESOLUTION_CACHE[cache_key] = str(best) if best else None
    return RESOURCE_RESOLUTION_CACHE[cache_key]


def canvas_assignment_url(config, entry):
    course_links = config.get("canvas_assignment_links", {}).get(entry["course"], {})
    if entry["category"] == "练习":
        return course_links.get("practice", {}).get(entry["unit"])
    if entry["category"] == "测试":
        filename = entry["filename"].lower()
        if "midterm" in filename:
            return course_links.get("test", {}).get("Midterm")
        if "final" in filename:
            return course_links.get("test", {}).get("Final")
    return None


def canvas_assignment_metadata(config, entry):
    course_links = config.get("canvas_assignment_links", {}).get(entry["course"], {})
    key = None
    if entry["category"] == "练习":
        key = entry["unit"]
    elif entry["category"] == "测试":
        filename = entry["filename"].lower()
        if "midterm" in filename:
            key = "Midterm"
        elif "final" in filename:
            key = "Final"
    meta = course_links.get("metadata", {}).get(key, {})
    return meta if isinstance(meta, dict) else {}


def canvas_assignment_inferred(config, entry):
    inferred = config.get("canvas_assignment_links", {}).get(entry["course"], {}).get("inferred", [])
    target = canvas_assignment_url(config, entry)
    return bool(target and target in inferred)


RESOURCE_LABELS = {
    "课本": "textbook",
    "教材": "textbook",
    "大纲": "syllabus",
    "课件": "courseware",
    "练习": "practice",
    "测试": "test",
    "真题": "frq",
    "模考": "mock_exam",
    "词汇": "vocabulary",
    "公式表": "formula_sheet",
    "刷题": "checklist",
}


def task_needs_practice(title, kind=None):
    text = f"{title or ''} {kind or ''}"
    return any(term in text for term in ["练习", "刷题", "Practice", "FRQ", "题", "测试", "模考", "Midterm", "Final"])


def resource_index_resources(config, course, unit, title="", kind=None):
    include_practice = task_needs_practice(title, kind)
    resources = []
    for entry in resource_index_entries(config, course):
        category = entry["category"]
        if not unit_matches(entry["unit"], unit):
            continue
        if category in {"真题", "模考"} and not task_needs_practice(title, kind):
            continue
        if category == "测试" and not any(term in f"{title} {kind}" for term in ["测试", "模考", "Midterm", "Final"]):
            continue
        if category == "练习" and course == "AP_Calculus_BC" and not include_practice:
            continue
        target = resolve_index_resource(config, entry)
        label = f"resource_index_{RESOURCE_LABELS.get(category, slug(category))}"
        canvas_url = canvas_assignment_url(config, entry) if not target else None
        if canvas_url:
            label = f"{label}_canvas"
            target = canvas_url
        resource = {
            "label": label,
            "target": target,
            "source": "canvas_assignment_links" if canvas_url else "workbook_resource_index",
            "index_filename": entry["filename"],
            "index_category": category,
            "index_unit": entry["unit"],
            "index_row": entry["row"],
        }
        if canvas_url and canvas_assignment_inferred(config, entry):
            resource["inferred"] = True
        if canvas_url:
            resource.update(canvas_assignment_metadata(config, entry))
        if target:
            resources.append(resource)
        else:
            resources.append({**resource, "missing": True, "target": ""})
    return resources


def supplemental_courseware_resources(config, course, unit):
    target = config.get("supplemental_courseware", {}).get(course, {}).get(str(unit))
    if target and Path(target).exists():
        return [{
            "label": "supplemental_courseware",
            "target": target,
            "source": "supplemental_courseware",
            "index_unit": str(unit),
        }]
    return []


def first_resource_target(resources, labels):
    wanted = set(labels)
    for resource in resources:
        if resource.get("label") in wanted:
            target = resource.get("target")
            if target and Path(str(target)).exists():
                return target
    return None


def parse_numbered_topic_range(title):
    text = str(title or "")
    match = re.search(r"(?<!\d)(\d+)\.(\d+)\s*-\s*(?:(\d+)\.)?(\d+)(?!\d)", text)
    if match:
        chapter = int(match.group(1))
        start = f"{chapter}.{int(match.group(2))}"
        end_chapter = int(match.group(3)) if match.group(3) else chapter
        end = f"{end_chapter}.{int(match.group(4))}"
        return start, end
    match = re.search(r"(?<!\d)(\d+)\.(\d+)(?!\d)", text)
    if match:
        section = f"{int(match.group(1))}.{int(match.group(2))}"
        return section, section
    return None


def increment_numbered_section(section):
    match = re.match(r"^(\d+)\.(\d+)$", str(section))
    if not match:
        return None
    return f"{int(match.group(1))}.{int(match.group(2)) + 1}"


def first_heading_page(index, section, start_after=1):
    prefix = f"{section} "
    for page in index.get("pages", []):
        page_no = page.get("page", 0)
        if page_no < start_after:
            continue
        text = (page.get("text") or "").strip()
        if text.startswith(prefix) or text == section:
            return page_no
    return None


def numbered_courseware_page_range(index, title):
    section_range = parse_numbered_topic_range(title)
    if not section_range:
        return None
    start_section, end_section = section_range
    start_page = first_heading_page(index, start_section, start_after=2)
    if not start_page:
        return None
    next_section = increment_numbered_section(end_section)
    next_page = first_heading_page(index, next_section, start_after=start_page + 1) if next_section else None
    pages = index.get("pages", [])
    last_page = pages[-1]["page"] if pages else start_page
    end_page = (next_page - 1) if next_page else min(last_page, start_page + 40)
    if end_page < start_page:
        end_page = start_page
    return [start_page, end_page], [("section_boundary", start_section), ("section_boundary", end_section)]


def bc_semantic_courseware_page_range(unit, title):
    unit_text = str(unit or "")
    title_text = str(title or "")
    if unit_text != "Unit 2":
        return None
    mappings = [
        (["导数定义", "切线斜率"], [3, 16], "2.1-2.2 derivative definition and tangent slope"),
        (["导数作为函数", "可微性"], [13, 25], "2.2-2.4 derivative as function and differentiability"),
        (["基本求导", "幂", "积", "商"], [26, 41], "2.5-2.9 basic derivative rules"),
        (["三角函数"], [34, 44], "2.7 and 2.10 trigonometric derivatives"),
    ]
    for triggers, page_range, reason in mappings:
        if any(trigger in title_text for trigger in triggers):
            return page_range, [("semantic_topic", reason)]
    return None


def courseware_excerpt_resources(config, course, unit, title, resources, row=None):
    source = first_resource_target(resources, ["resource_index_courseware", "primary_resource", "supplemental_courseware"])
    if not source:
        return []
    try:
        index_key = f"courseware_{course}_{slug(unit)}_{hashlib.sha1(str(source).encode('utf-8')).hexdigest()[:8]}"
        index = pdf_text_index(source, index_key)
    except Exception:
        return []
    keywords = topic_print_keywords(course, unit, title, row)
    max_pages = 10 if course == "AP_Calculus_BC" else 14
    result = bc_semantic_courseware_page_range(unit, title) if course == "AP_Calculus_BC" else None
    match_method = "semantic_topic" if result else "topic_keywords"
    if not result:
        result = numbered_courseware_page_range(index, title) if index else None
        match_method = "section_boundary" if result else "topic_keywords"
    if not result:
        result = best_keyword_window(index, keywords, max_pages=max_pages, context_pages=1) if index else None
    if result:
        page_range, hits = result
        start_page, end_page = page_range
    else:
        start_page, end_page = 1, min(max_pages, len(index.get("pages", [])) or max_pages)
        hits = []
        match_method = "courseware_start_fallback"
    excerpt = create_pdf_excerpt(source, start_page, end_page, f"{course} {unit} courseware {title}")
    if not excerpt:
        return []
    return [{
        "label": "task_excerpt_courseware",
        "target": excerpt,
        "source": source,
        "page_range": [start_page, end_page],
        "match": match_method,
        "top_hits": hits,
    }]


def dedupe_resources(resources):
    seen = set()
    clean = []
    for resource in resources:
        target = resource.get("target")
        if not target and resource.get("missing"):
            target = f"missing:{resource.get('label')}:{resource.get('index_filename')}"
        if not target or target in seen:
            continue
        seen.add(target)
        clean.append(resource)
    return clean


def pdf_excerpt_path(source_path, start_page, end_page, label):
    source = Path(source_path)
    digest = hashlib.sha1(f"{source}:{start_page}:{end_page}:{label}".encode("utf-8")).hexdigest()[:10]
    safe_label = slug(label)
    return TASK_MATERIALS_DIR / f"{safe_label}_p{start_page}-{end_page}_{digest}.pdf"


def create_pdf_excerpt(source_path, start_page, end_page, label):
    source = Path(source_path)
    if not source.exists():
        return None
    output = pdf_excerpt_path(source, start_page, end_page, label)
    if output.exists() and output.stat().st_size > 0:
        return str(output)
    TASK_MATERIALS_DIR.mkdir(parents=True, exist_ok=True)
    from pypdf import PdfReader, PdfWriter
    try:
        reader = PdfReader(str(source))
        total = len(reader.pages)
        start = max(1, min(int(start_page), total))
        end = max(start, min(int(end_page), total))
        writer = PdfWriter()
        for page_index in range(start - 1, end):
            writer.add_page(reader.pages[page_index], excluded_keys=["/Annots", "/B", "/StructParents"])
    except Exception as exc:
        warning_path = TASK_MATERIALS_DIR / "excerpt_errors.log"
        with warning_path.open("a", encoding="utf-8") as f:
            f.write(f"{dt.datetime.now().isoformat(timespec='seconds')} failed {source} p{start_page}-{end_page}: {exc}\n")
        return None
    metadata = {
        "/Title": f"AP Learning OS - {label}",
        "/Subject": f"{source.name} pages {start}-{end}"
    }
    try:
        writer.add_metadata(metadata)
    except Exception:
        pass
    with output.open("wb") as f:
        writer.write(f)
    return str(output)


def flatten_outline_items(outline):
    for item in outline:
        if isinstance(item, list):
            yield from flatten_outline_items(item)
        else:
            yield item


def pdf_section_index(source_path, cache_key):
    cache_path = BASE / "data" / f"{cache_key}_section_index.json"
    source = Path(source_path)
    if cache_path.exists():
        cached = read_json(cache_path)
        if cached.get("source_mtime") == source.stat().st_mtime:
            return cached["sections"]
    from pypdf import PdfReader
    reader = PdfReader(str(source))
    sections = {}
    for item in flatten_outline_items(reader.outline):
        title = getattr(item, "title", str(item)).replace("\x00", "")
        m = re.search(r"\b(\d+)\.(\d+)\b", title)
        if not m:
            continue
        try:
            page = reader.get_destination_page_number(item) + 1
        except Exception:
            continue
        sections[f"{m.group(1)}.{m.group(2)}"] = page
    write_json(cache_path, {"source": str(source), "source_mtime": source.stat().st_mtime, "sections": sections})
    return sections


def pdf_text_index(source_path, cache_key, max_chars=4000):
    source = Path(source_path)
    if not source.exists():
        return None
    cache_path = BASE / "data" / f"{cache_key}_text_index.json"
    if cache_path.exists():
        cached = read_json(cache_path)
        if cached.get("source_mtime") == source.stat().st_mtime:
            return cached
    from pypdf import PdfReader
    pages = []
    with open(os.devnull, "w") as devnull:
        with contextlib.redirect_stderr(devnull):
            reader = PdfReader(str(source))
            for idx, page in enumerate(reader.pages, start=1):
                try:
                    text = page.extract_text() or ""
                except Exception:
                    text = ""
                pages.append({"page": idx, "text": " ".join(text.split())[:max_chars]})
    index = {"source": str(source), "source_mtime": source.stat().st_mtime, "pages": pages}
    write_json(cache_path, index)
    return index


def next_section_page(section_pages, section):
    parts = section.split(".")
    if len(parts) != 2:
        return None
    ch, sec = int(parts[0]), int(parts[1])
    candidates = []
    for key, page in section_pages.items():
        m = re.match(r"^(\d+)\.(\d+)$", key)
        if not m:
            continue
        kch, ksec = int(m.group(1)), int(m.group(2))
        if (kch, ksec) > (ch, sec):
            candidates.append(page)
    return min(candidates) if candidates else None


def parse_stewart_sections(text):
    if not text or "Stewart" not in str(text):
        return []
    raw = str(text)
    matches = []
    pattern = re.compile(r"(\d+)\.(\d+)(?:\s*-\s*(?:(\d+)\.)?(\d+))?")
    for m in pattern.finditer(raw):
        start = f"{m.group(1)}.{m.group(2)}"
        if m.group(4):
            end_ch = m.group(3) or m.group(1)
            end = f"{end_ch}.{m.group(4)}"
        else:
            end = start
        matches.append((start, end))
    return matches


def stewart_excerpt_resources(config, row):
    material = config.get("canonical_materials", {}).get("AP_Calculus_BC", {}).get("stewart_textbook")
    if not material or not Path(material).exists():
        return []
    refs = parse_stewart_sections(row.get("课件/课本参考"))
    if not refs:
        return []
    section_pages = pdf_section_index(material, "stewart")
    resources = []
    for start_section, end_section in refs:
        start_page = section_pages.get(start_section)
        end_start_page = section_pages.get(end_section)
        if not start_page or not end_start_page:
            continue
        next_page = next_section_page(section_pages, end_section)
        end_page = (next_page - 1) if next_page else (end_start_page + 12)
        label = f"BC {row.get('Unit')} {row.get('天数')} Stewart {start_section}-{end_section}"
        excerpt = create_pdf_excerpt(material, start_page, end_page, label)
        if excerpt:
            resources.append({
                "label": "task_excerpt_stewart",
                "target": excerpt,
                "source": material,
                "page_range": [start_page, end_page],
                "sections": [start_section, end_section]
            })
    resources.append({"label": "canonical_stewart_textbook", "target": material})
    master = config.get("canonical_materials", {}).get("AP_Calculus_BC", {}).get("ap_master_packet")
    if master and Path(master).exists():
        resources.append({"label": "canonical_ap_master_packet", "target": master})
    return resources


def csa_textbook_keywords(unit, title):
    text = str(title or "")
    unit_text = str(unit or "")
    keywords = [
        ("Java", 2),
        ("program", 2),
        ("chapter contents", -8),
        ("chapter summary", -8),
        ("skill practice", -8),
        ("multiple choice exercises", -8),
        ("index", -5),
    ]
    if "Unit 1" in unit_text:
        keywords.extend([
            ("Programming Building Blocks", 8),
            ("Java Basics", 8),
            ("program has two elements", 12),
            ("instructions and data", 12),
            ("input the data", 10),
            ("output the results", 10),
            ("data types variables and constants", 14),
            ("declaring variables", 12),
            ("integer data types", 10),
            ("floating-point data types", 10),
            ("boolean data type", 8),
            ("assignment operator", 10),
            ("expressions and arithmetic operators", 12),
            ("String literals", 8),
            ("Java Application Structure", 2),
        ])
    concept_map = {
        "算法": [("algorithm", 10), ("instructions", 8), ("processing", 8), ("input the data", 8), ("output the results", 8)],
        "变量": [("variable", 10), ("variables", 10), ("declaring variables", 12), ("named locations in memory", 10)],
        "数据类型": [("data type", 10), ("data types", 10), ("primitive", 8), ("int", 3), ("double", 3), ("boolean", 3), ("char", 3)],
        "表达式": [("expression", 10), ("expressions", 10), ("arithmetic operators", 10), ("operator precedence", 8)],
        "输出": [("output", 10), ("print", 6), ("println", 6), ("System.out", 8)],
        "赋值": [("assignment operator", 12), ("initial values", 8), ("literals", 6)],
        "输入": [("input", 10), ("keyboard", 6), ("Scanner", 12)],
        "Scanner": [("Scanner", 14), ("keyboard input", 10), ("java.util.Scanner", 12)],
        "类型转换": [("type conversion", 14), ("casting", 10), ("compatible data types", 8)],
        "复合赋值": [("compound assignment", 14), ("increment", 8), ("decrement", 8)],
        "API": [("API", 10), ("library", 8), ("documentation", 8)],
        "库": [("library", 10), ("package", 6), ("import", 6)],
        "注释": [("comment", 10), ("comments", 10), ("documentation", 6)],
        "方法": [("method", 8), ("methods", 8), ("method signature", 12), ("parameters", 8), ("arguments", 8)],
        "调用": [("calling methods", 12), ("method call", 12), ("arguments", 8)],
        "Math": [("Math class", 14), ("Math.", 10)],
        "对象": [("object", 8), ("objects", 8), ("instantiation", 8)],
        "String": [("String class", 14), ("substring", 10), ("indexOf", 10), ("equals", 10), ("length", 8)],
    }
    for trigger, mapped in concept_map.items():
        if trigger in text:
            keywords.extend(mapped)
    for token in re.findall(r"[A-Za-z][A-Za-z0-9_.]{2,}", text):
        keywords.append((token, 3))
    return keywords


def best_keyword_window(index, keywords, max_pages=28, context_pages=1):
    pages = index.get("pages", [])
    if not pages:
        return None
    scored = [(page["page"], keyword_score(page.get("text", ""), keywords)) for page in pages]
    best = None
    left = 0
    running = 0
    for right, (_, score) in enumerate(scored):
        running += score
        while scored[right][0] - scored[left][0] + 1 > max_pages:
            running -= scored[left][1]
            left += 1
        if best is None or running > best[0]:
            best = (running, left, right)
    if not best or best[0] <= 0:
        return None
    _, left, right = best
    while left <= right and scored[left][1] <= 0:
        left += 1
    while right >= left and scored[right][1] <= 0:
        right -= 1
    if left > right:
        return None
    start = max(pages[0]["page"], scored[left][0] - context_pages)
    end = min(pages[-1]["page"], scored[right][0] + context_pages)
    hits = sorted([(score, page) for page, score in scored if start <= page <= end and score > 0], reverse=True)[:8]
    return [start, end], hits


def csa_excerpt_resources(config, unit, title, include_practice=False):
    material = config.get("canonical_materials", {}).get("AP_CSA", {}).get("java_textbook")
    resources = []
    if material and Path(material).exists():
        index = java_textbook_scope(pdf_text_index(material, "java_illuminated"), unit)
        result = best_keyword_window(index, csa_textbook_keywords(unit, title), max_pages=32, context_pages=1) if index else None
        if result:
            page_range, hits = result
            start_page, end_page = page_range
        else:
            start_page, end_page = 1, 12
            hits = []
        excerpt = create_pdf_excerpt(material, start_page, end_page, f"CSA {unit} Java Illuminated {title}")
        if excerpt:
            resources.append({
                "label": "task_excerpt_java_illuminated",
                "target": excerpt,
                "source": material,
                "page_range": [start_page, end_page],
                "match": "topic_keywords",
                "top_hits": hits,
            })
        resources.append({"label": "canonical_java_textbook", "target": material})
    master = config.get("canonical_materials", {}).get("AP_CSA", {}).get("ap_master_packet")
    if master and Path(master).exists():
        resources.append({"label": "canonical_ap_master_packet", "target": master})
    return resources


def print_packet_text_index(config):
    packet = config.get("dated_print_packet", {})
    pdf = packet.get("pdf")
    if not pdf:
        return None
    return pdf_text_index(pdf, "dated_print_packet")


def keyword_score(text, keywords):
    lowered = text.lower()
    score = 0
    for keyword, weight in keywords:
        if not keyword:
            continue
        occurrences = lowered.count(keyword.lower())
        score += occurrences * weight
    return score


def scoped_index(index, start_page=None, end_page=None):
    if not index or (start_page is None and end_page is None):
        return index
    pages = []
    for page in index.get("pages", []):
        number = page["page"]
        if start_page is not None and number < start_page:
            continue
        if end_page is not None and number > end_page:
            continue
        pages.append(page)
    if not pages:
        return index
    scoped = dict(index)
    scoped["pages"] = pages
    scoped["scope"] = [pages[0]["page"], pages[-1]["page"]]
    return scoped


def unit_number(unit):
    match = re.search(r"\d+", str(unit or ""))
    return int(match.group(0)) if match else None


def find_text_page(index, patterns, start_after=0):
    lowered_patterns = [pattern.lower() for pattern in patterns if pattern]
    for page in index.get("pages", []):
        if page["page"] <= start_after:
            continue
        text = page.get("text", "").lower()
        if any(pattern in text for pattern in lowered_patterns):
            return page["page"]
    return None


def csa_print_packet_scope(index, unit):
    number = unit_number(unit)
    if not number:
        return index
    start = find_text_page(index, [f"AP COMPUTER SCIENCE A UNIT {number}"])
    if not start:
        return index
    next_start = find_text_page(index, [f"AP COMPUTER SCIENCE A UNIT {number + 1}"], start_after=start)
    end = next_start - 1 if next_start else start + 120
    return scoped_index(index, start, end)


def java_textbook_scope(index, unit):
    if str(unit or "") == "Unit 1":
        start = find_text_page(index, ["CHAPTER 2 Programming Building Blocks", "Programming Building Blocks—Java Basics"], start_after=40)
        if start:
            next_start = find_text_page(index, ["CHAPTER 3"], start_after=start)
            return scoped_index(index, start, (next_start - 1) if next_start else start + 90)
    return index


def topic_print_keywords(course, unit, title, row=None):
    title = str(title or "")
    unit = str(unit or "")
    keywords = []
    if course == "AP_Calculus_BC":
        keywords.extend([
            ("AP Calculus BC", 1),
            ("CALCULUS BC", 1),
            ("derivative", 5),
            ("differentiation", 4),
            ("tangent", 3),
            ("rate of change", 4),
            ("导数", 5),
            ("切线", 4),
            ("SCORING GUIDELINES", -40),
            ("Question ", -8),
            ("Free Response", -12),
        ])
        if row:
            refs = parse_stewart_sections(row.get("课件/课本参考"))
            for start, end in refs:
                keywords.append((start, 2))
                if end != start:
                    keywords.append((end, 2))
        if "Unit 1" in unit or "极限" in title:
            keywords.extend([("limit", 5), ("continuity", 4), ("horizontal asymptote", 3), ("极限", 5)])
        if "Unit 2" in unit or "导数" in title:
            keywords.extend([
                ("Defining the Derivative of a Function", 80),
                ("defining the derivative", 60),
                ("definition of derivative", 55),
                ("derivative notation", 40),
                ("tangent line", 25),
                ("basic derivative", 6),
                ("Topic 2.D", 45),
                ("slope field", -12),
            ])
        if "Unit 3" in unit:
            keywords.extend([("chain rule", 7), ("implicit", 5), ("inverse", 4)])
        if "Unit 4" in unit:
            keywords.extend([("related rates", 6), ("motion", 5), ("linear approximation", 4)])
    elif course == "AP_CSA":
        keywords.extend([
            ("AP Computer Science A", 5),
            ("Computer Science A", 4),
            ("Java", 3),
        ])
        if "Unit 1" in unit:
            keywords.extend([
                ("UNIT 1", 2),
                ("Unit 1: Using Objects and Methods", 4),
                ("Using Objects", 3),
                ("Chapter Introduction and Learning Strategy", 3),
                ("Introduction to Java", 4),
                ("Primitive data type", 6),
                ("primitive types", 6),
                ("data types", 6),
                ("expressions", 5),
                ("output", 5),
                ("main method", 4),
                ("Methods", 4),
                ("variable", 3),
                ("expression", 4),
                ("data type", 4),
                ("assignment", 4),
                ("String", 3),
                ("变量", 5),
                ("表达式", 5),
                ("UNIT 5", -30),
                ("Scope and Access", -20),
                ("Global and Local Variables", -20),
                ("constructor", -6),
            ])
            if any(k in title for k in ["1.1", "1.2", "1.3", "算法", "数据类型", "表达式", "输出"]):
                keywords.extend([
                    ("Chapter Introduction and Learning Strategy", 18),
                    ("Introduction to Java", 16),
                    ("Primitive data type", 14),
                    ("创建变量", 10),
                    ("Output Code", 12),
                    ("输出语句", 12),
                    ("Operator 运算符", 8),
                ])
            if any(k in title for k in ["1.4", "赋值", "输入", "Scanner"]):
                keywords.extend([
                    ("赋值操作", 20),
                    ("Other Methods to Create a Variable", 18),
                    ("其他创建变量的方式", 18),
                    ("Change Value", 16),
                    ("更改变量", 16),
                    ("Output Code", 14),
                    ("输出语句", 14),
                    ("System.out.println", 12),
                    ("Scanner", 20),
                    ("输入", 16),
                ])
        if "Unit 2" in unit:
            keywords.extend([("boolean", 5), ("if", 3), ("loop", 5), ("iteration", 5), ("while", 4), ("for", 4)])
        if "Unit 3" in unit:
            keywords.extend([("class", 5), ("constructor", 5), ("object", 4), ("method", 4)])
        if "刷题" in title:
            keywords.extend([("Test Booklet", 5), ("Scoring Guide", 3)])
    for token in re.findall(r"[A-Za-z][A-Za-z0-9_]{2,}|[\u4e00-\u9fff]{2,}", title):
        keywords.append((token, 2))
    return keywords


def best_topic_page_range(index, keywords, max_pages=10):
    return best_keyword_window(index, keywords, max_pages=max_pages, context_pages=0)


def best_csa_print_packet_page_range(index, keywords):
    return best_keyword_window(index, keywords, max_pages=10, context_pages=4)


def topic_print_packet_resources(config, course, unit, title, row=None):
    packet = config.get("dated_print_packet", {})
    index = print_packet_text_index(config)
    if not index:
        return []
    if course == "AP_CSA":
        index = csa_print_packet_scope(index, unit)
    keywords = topic_print_keywords(course, unit, title, row)
    result = best_csa_print_packet_page_range(index, keywords) if course == "AP_CSA" else best_topic_page_range(index, keywords)
    if not result:
        return []
    page_range, hits = result
    match = "topic_keywords"
    start_page, end_page = page_range
    pdf = index["source"]
    label = f"{course} topic print packet {unit} {title}"
    excerpt = create_pdf_excerpt(pdf, start_page, end_page, label)
    if not excerpt:
        return []
    excerpt_text = ""
    try:
        selected_pages = [
            page for page in index.get("pages", [])
            if start_page <= page.get("page", 0) <= end_page
        ]
        excerpt_text = " ".join(page.get("text", "") for page in selected_pages).lower()
    except Exception:
        excerpt_text = ""
    bad_markers = [
        "practice exam",
        "test booklet",
        "scoring guide",
        "scoring guidelines",
        "free-response questions",
        "answer key",
        "practice book",
    ]
    if any(marker in excerpt_text for marker in bad_markers):
        warning_path = TASK_MATERIALS_DIR / "material_match_warnings.log"
        with warning_path.open("a", encoding="utf-8") as f:
            f.write(
                f"{dt.datetime.now().isoformat(timespec='seconds')} rejected {course} {unit} {title} "
                f"p{start_page}-{end_page}: exam/scoring marker detected\n"
            )
        return []
    return [{
        "label": f"supplemental_print_packet_{'bc' if course == 'AP_Calculus_BC' else 'csa'}",
        "target": excerpt,
        "source": pdf,
        "page_range": [start_page, end_page],
        "match": match,
        "top_hits": hits,
        "note": packet.get("note", "")
    }]

def build_csa_tasks(config, state, current_date):
    plan = Path(config["plans"]["AP_CSA"])
    wb = load_workbook(plan, data_only=True, read_only=True)
    ws = wb["每日刷题计划"]
    tasks = []
    planner = config["planner"]
    csa_date_offset = dt.timedelta(0)
    if planner.get("csa_plan_start_date") and planner.get("csa_actual_start_date"):
        csa_date_offset = parse_date_arg(planner["csa_actual_start_date"]) - parse_date_arg(planner["csa_plan_start_date"])
    for row in row_dicts(ws, 3):
        row_date = normalize_plan_mmdd(row.get("日期"), planner["csa_year"])
        if row_date:
            row_date = row_date + csa_date_offset
        if row_date != current_date:
            continue
        practice_ids = str(row.get("今日刷题编号") or "")
        is_practice = "—" not in practice_ids and practice_ids.strip() not in {"", "-"}
        kind = "CSA_PRACTICE" if is_practice else "CSA_CONCEPT"
        default_min = config["planner"]["default_targets_min"][kind]
        task_key = f"CSA:{row.get('Unit')}:{row.get('天数')}:{kind}"
        target_min = adaptive_minutes(state, task_key, default_min, config)
        resources = []
        for label, col in [("question_file", "题目文件"), ("answer_file", "答案文件")]:
            found = find_existing_path(config, row.get(col))
            if found:
                resources.append({"label": label, "target": found})
        resources.extend(resource_index_resources(
            config,
            "AP_CSA",
            row.get("Unit"),
            row.get("学习内容"),
            kind,
        ))
        resources.extend(supplemental_courseware_resources(config, "AP_CSA", row.get("Unit")))
        resources.extend(courseware_excerpt_resources(
            config,
            "AP_CSA",
            row.get("Unit"),
            row.get("学习内容"),
            resources,
            row,
        ))
        resources.extend(csa_excerpt_resources(config, row.get("Unit"), row.get("学习内容"), include_practice=is_practice))
        resources = dedupe_resources(resources)
        tasks.append({
            "id": f"CSA-{current_date.isoformat()}-{slug(row.get('天数'))}-{kind}",
            "task_key": task_key,
            "course": "AP_CSA",
            "kind": kind,
            "state": "Planned",
            "date": current_date.isoformat(),
            "unit": row.get("Unit"),
            "title": f"CSA {row.get('Unit')} {row.get('天数')} - {row.get('学习内容')}",
            "target_min": target_min,
            "observable_goal": (
                f"Open CSA materials for {row.get('学习内容')} and work for at least {target_min} minutes."
            ),
            "completion_criteria": [
                f"CSA-related resource/window remains active for about {target_min} minutes",
                "Start/end screenshots exist",
                "uone reviews evidence and overwrites final status"
            ],
            "resources": resources,
            "launch": [
                {"type": "url", "target": config["urls"]["khan_ap_csa"]},
                {"type": "app", "target": config["apps"]["notes"]},
                {"type": "app", "target": config["apps"]["editor"]}
            ],
            "source": {"workbook": str(plan), "sheet": "每日刷题计划", "row": row},
        })
    return tasks


def build_bc_tasks(config, state, current_date):
    plan = Path(config["plans"]["AP_Calculus_BC"])
    wb = load_workbook(plan, data_only=True, read_only=True)
    ws = wb["每日计划"]
    bc_rows = []
    cumulative_day = 0
    for row in row_dicts(ws, 3):
        rng = parse_day_range(row.get("天数"))
        if not rng:
            continue
        start, end = rng
        duration_days = max(1, end - start + 1)
        global_start = cumulative_day + 1
        global_end = cumulative_day + duration_days
        cumulative_day = global_end
        bc_rows.append((row, global_start, global_end))
    anchor = config.get("planner", {}).get("bc_anchor")
    if anchor:
        anchor_date = parse_date_arg(anchor["date"])
        anchor_global_day = None
        for row, global_start, global_end in bc_rows:
            if str(row.get("Unit")) == str(anchor.get("unit")) and str(row.get("天数")) == str(anchor.get("day")):
                anchor_global_day = global_start
                break
        if anchor_global_day is None:
            raise SystemExit(f"BC anchor not found: {anchor}")
        global_day_number = anchor_global_day + (current_date - anchor_date).days
    else:
        day1 = parse_date_arg(config["planner"]["bc_day_1_date"])
        global_day_number = (current_date - day1).days + 1
    if global_day_number < 1:
        return []
    tasks = []
    for row, global_start, global_end in bc_rows:
        if not (global_start <= global_day_number <= global_end):
            continue
        task_key = f"BC:{row.get('Unit')}:{row.get('天数')}:RESOURCE"
        default_min = config["planner"]["default_targets_min"]["BC_RESOURCE_WORK"]
        target_min = adaptive_minutes(state, task_key, default_min, config)
        primary = find_existing_path(config, row.get("文件名（可直接打开）"))
        resources = []
        if primary:
            resources.append({"label": "primary_resource", "target": primary})
        for label, col in [("reference", "课件/课本参考"), ("practice", "配套练习/答案")]:
            target = find_existing_path(config, row.get(col))
            if target:
                resources.append({"label": label, "target": target})
        resources.extend(resource_index_resources(
            config,
            "AP_Calculus_BC",
            row.get("Unit"),
            row.get("学习内容"),
            "BC_RESOURCE_WORK",
        ))
        resources.extend(courseware_excerpt_resources(
            config,
            "AP_Calculus_BC",
            row.get("Unit"),
            row.get("学习内容"),
            resources,
            row,
        ))
        resources.extend(stewart_excerpt_resources(config, row))
        resources = dedupe_resources(resources)
        launch_resource = next(
            (r.get("target") for r in resources if r.get("label") == "task_excerpt_courseware"),
            None
        ) or primary or next(
            (r.get("target") for r in resources if str(r.get("label", "")).startswith("resource_index_")),
            None
        )
        tasks.append({
            "id": f"BC-{current_date.isoformat()}-DAY{global_day_number}-{slug(row.get('Unit'))}",
            "task_key": task_key,
            "course": "AP_Calculus_BC",
            "kind": "BC_RESOURCE_WORK",
            "state": "Planned",
            "date": current_date.isoformat(),
            "unit": row.get("Unit"),
            "title": f"BC {row.get('Unit')} {row.get('天数')} - {row.get('学习内容')}",
            "target_min": target_min,
            "observable_goal": (
                f"Open BC resource for {row.get('学习内容')} and work for at least {target_min} minutes."
            ),
            "completion_criteria": [
                f"BC resource/Khan/GoodNotes window evidence for about {target_min} minutes",
                "Start/end screenshots exist",
                "uone reviews evidence and overwrites final status"
            ],
            "resources": resources,
            "launch": [
                {"type": "url", "target": config["urls"].get("khan_calculus_bc_units", {}).get(str(row.get("Unit")), config["urls"]["khan_calculus_bc"])},
                {"type": "app", "target": config["apps"]["notes"]},
                {"type": "resource", "target": launch_resource} if launch_resource else {"type": "app", "target": config["apps"]["pdf"]}
            ],
            "source": {
                "workbook": str(plan),
                "sheet": "每日计划",
                "row": row,
                "mapped_global_day_number": global_day_number,
                "mapped_global_day_range": [global_start, global_end]
            },
        })
    return tasks


def add_error_log_task(config, state, current_date):
    key = f"ERROR_LOG:{current_date.isoformat()}"
    default_min = config["planner"]["default_targets_min"]["CSA_ERROR_LOG"]
    target_min = adaptive_minutes(state, key, default_min, config)
    return {
        "id": f"ERRORLOG-{current_date.isoformat()}",
        "task_key": key,
        "course": "GLOBAL",
        "kind": "ERROR_LOG",
        "state": "Planned",
        "date": current_date.isoformat(),
        "unit": "All",
        "title": "Error Log Review",
        "target_min": target_min,
        "observable_goal": f"Open the error log/notes and review mistakes for at least {target_min} minutes.",
        "completion_criteria": [
            f"GoodNotes or notes window active for about {target_min} minutes",
            "Review screenshot exists",
            "uone reviews evidence and overwrites final status"
        ],
        "resources": [],
        "launch": [{"type": "app", "target": config["apps"]["notes"]}],
        "source": {"generated": True}
    }


def generate_today(config, state, current_date):
    tasks = []
    tasks.extend(build_bc_tasks(config, state, current_date))
    tasks.extend(build_csa_tasks(config, state, current_date))
    if tasks:
        tasks.append(add_error_log_task(config, state, current_date))
    generated_ids = {task["id"] for task in tasks}
    for task_id, existing in list(state["tasks"].items()):
        if existing.get("date") == current_date.isoformat() and task_id not in generated_ids:
            if existing.get("state") in {"Planned", "Needs Review", "Failed", "Partially Completed"}:
                del state["tasks"][task_id]
    for task in tasks:
        existing = state["tasks"].get(task["id"], {})
        task["state"] = existing.get("state", task["state"])
        task["last_session_id"] = existing.get("last_session_id")
        state["tasks"][task["id"]] = task
    save_state(state)
    write_blackboard(tasks, current_date)
    return tasks


def write_blackboard(tasks, current_date):
    lines = [
        f"# AP Learning OS Blackboard - {current_date.isoformat()}",
        "",
        "Rule: detectors produce evidence only. Final completion requires uone review.",
        "",
    ]
    if not tasks:
        lines.append("No planned tasks found for this date.")
    for task in tasks:
        lines.extend([
            f"## [{task['state']}] {task['id']}",
            "",
            f"- Course: {task['course']}",
            f"- Unit: {task['unit']}",
            f"- Target time: {task['target_min']} min",
            f"- Observable goal: {task['observable_goal']}",
            "- Completion criteria:",
        ])
        for criterion in task["completion_criteria"]:
            lines.append(f"  - {criterion}")
        if task.get("resources"):
            lines.append("- Resources:")
            for resource in task["resources"]:
                lines.append(f"  - {resource['label']}: {resource['target']}")
        lines.append("")
    BLACKBOARD_PATH.parent.mkdir(parents=True, exist_ok=True)
    BLACKBOARD_PATH.write_text("\n".join(lines), encoding="utf-8")


def run(cmd, check=False):
    return subprocess.run(cmd, text=True, capture_output=True, check=check)


def osascript(script):
    return run(["osascript", "-e", script]).stdout.strip()


def frontmost_window():
    script = '''
tell application "System Events"
  set frontApp to name of first application process whose frontmost is true
  set winTitle to ""
  try
    tell process frontApp
      if count of windows > 0 then set winTitle to name of front window
    end tell
  end try
  return frontApp & " | " & winTitle
end tell
'''
    try:
        value = osascript(script)
    except Exception as exc:
        value = f"UNKNOWN | {exc}"
    app, _, title = value.partition(" | ")
    return {"app": app, "window": title, "timestamp": dt.datetime.now().isoformat(timespec="seconds")}


def capture_screenshot(path):
    path.parent.mkdir(parents=True, exist_ok=True)
    result = run(["screencapture", "-x", str(path)])
    ok = result.returncode == 0 and path.exists()
    if not ok:
        error_path = path.parent / "screenshot_errors.log"
        error_path.write_text(
            f"{dt.datetime.now().isoformat(timespec='seconds')} failed to capture {path.name}\n"
            f"returncode={result.returncode}\nstdout={result.stdout}\nstderr={result.stderr}\n"
            "macOS may need Screen Recording permission for the terminal/Codex process.\n",
            encoding="utf-8"
        )
    return ok


def launch_item(item):
    typ = item.get("type")
    target = item.get("target")
    if not target:
        return
    if typ == "url":
        subprocess.Popen(["open", target])
    elif typ == "app":
        subprocess.Popen(["open", "-a", target])
    elif typ == "resource":
        if Path(target).exists():
            subprocess.Popen(["open", target])


def launch_task_resources(task):
    if task.get("course") == "AP_Calculus_BC":
        open_label_order = [
            "resource_index_practice_canvas",
            "resource_index_practice",
            "resource_index_test_canvas",
            "resource_index_test",
            "resource_index_frq",
            "resource_index_mock_exam",
            "task_excerpt_courseware",
        ]
    else:
        open_label_order = [
            "resource_index_courseware",
            "supplemental_courseware",
            "resource_index_syllabus",
            "resource_index_textbook",
            "resource_index_practice_canvas",
            "resource_index_practice",
            "resource_index_test_canvas",
            "resource_index_test",
            "resource_index_frq",
            "resource_index_mock_exam",
            "resource_index_checklist",
            "resource_index_vocabulary",
            "resource_index_formula_sheet",
            "question_file",
            "task_excerpt_courseware",
        ]
    resources_by_label = {}
    for resource in task.get("resources", []):
        resources_by_label.setdefault(resource.get("label"), []).append(resource)
    for label in open_label_order:
        for resource in resources_by_label.get(label, []):
            target = resource.get("target")
            if is_url(target):
                subprocess.Popen(["open", target])
            elif target and Path(target).exists():
                subprocess.Popen(["open", target])


def set_task_state(state, task_id, new_state):
    if new_state not in TASK_STATES:
        raise ValueError(f"invalid state: {new_state}")
    task = state["tasks"][task_id]
    old = task.get("state", "Planned")
    task["state"] = new_state
    task.setdefault("transitions", []).append({
        "from": old,
        "to": new_state,
        "timestamp": dt.datetime.now().isoformat(timespec="seconds")
    })


def compute_signal_score(task, samples, session_dir, observed_elapsed_min=None):
    target_min = float(task["target_min"])
    elapsed_min = float(observed_elapsed_min or 0)
    if elapsed_min <= 0 and samples:
        first = dt.datetime.fromisoformat(samples[0]["timestamp"])
        last = dt.datetime.fromisoformat(samples[-1]["timestamp"])
        elapsed_min = max(elapsed_min, (last - first).total_seconds() / 60)
    screenshots = list(session_dir.glob("*.png"))
    time_score = min(1.0, elapsed_min / target_min) if target_min else 0
    screenshot_score = min(1.0, len(screenshots) / 2)
    window_score = 0.0
    course_terms = {
        "AP_CSA": ["Safari", "Code", "GoodNotes", "Java", "Khan"],
        "AP_Calculus_BC": ["Safari", "Preview", "GoodNotes", "Khan", "Calculus"],
        "GLOBAL": ["GoodNotes", "Notes", "Preview"]
    }.get(task["course"], [])
    if samples and course_terms:
        hits = 0
        for sample in samples:
            haystack = f"{sample.get('app','')} {sample.get('window','')}".lower()
            if any(term.lower() in haystack for term in course_terms):
                hits += 1
        window_score = hits / len(samples)
    score = 0.5 * time_score + 0.25 * screenshot_score + 0.25 * window_score
    return {
        "score": round(score, 3),
        "elapsed_min_estimate": round(elapsed_min, 2),
        "screenshots": len(screenshots),
        "window_match_ratio": round(window_score, 3)
    }


def build_review_packet(session, task, signal_summary, session_dir):
    packet = session_dir / "review_packet.md"
    lines = [
        f"# uone Review Packet - {session['id']}",
        "",
        f"- Task: {task['id']}",
        f"- Title: {task['title']}",
        f"- Target: {task['target_min']} min",
        f"- System recommendation: {session['recommendation']}",
        f"- Signal score: {signal_summary['score']}",
        f"- Estimated elapsed: {signal_summary['elapsed_min_estimate']} min",
        f"- Window match ratio: {signal_summary['window_match_ratio']}",
        f"- Screenshots: {signal_summary['screenshots']}",
        "",
        "## Completion Criteria",
    ]
    for criterion in task["completion_criteria"]:
        lines.append(f"- {criterion}")
    lines.extend(["", "## Evidence Files"])
    for path in sorted(session_dir.iterdir()):
        if path.name != "review_packet.md":
            lines.append(f"- {path}")
    if (session_dir / "screenshot_errors.log").exists():
        lines.extend([
            "",
            "## Capture Warnings",
            "",
            "Screenshot capture failed. Grant Screen Recording permission to the terminal/Codex process in macOS System Settings, then rerun the session."
        ])
    lines.extend([
        "",
        "## uone Decision",
        "",
        "Run one of:",
        "",
        f"```bash",
        f"python3 {BASE / 'learning_os.py'} review --session {session['id']} --status completed --notes \"uone confirmed\"",
        f"python3 {BASE / 'learning_os.py'} review --session {session['id']} --status partially_completed --notes \"what remains\"",
        f"python3 {BASE / 'learning_os.py'} review --session {session['id']} --status not_completed --notes \"why\"",
        f"```",
    ])
    packet.write_text("\n".join(lines), encoding="utf-8")
    return str(packet)


def start_session(args):
    config = load_config()
    state = load_state()
    if args.task not in state["tasks"]:
        raise SystemExit(f"Unknown task id: {args.task}. Run `today` first.")
    task = state["tasks"][args.task]
    session_id = f"S-{dt.datetime.now().strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:6]}"
    session_dir = BASE / "sessions" / session_id
    session_dir.mkdir(parents=True, exist_ok=True)
    duration_min = args.duration_min if args.duration_min is not None else float(task["target_min"])
    set_task_state(state, args.task, "Running")
    task["last_session_id"] = session_id
    session = {
        "id": session_id,
        "task_id": args.task,
        "started_at": dt.datetime.now().isoformat(timespec="seconds"),
        "duration_min": duration_min,
        "state": "Running",
        "samples": [],
        "session_dir": str(session_dir)
    }
    state["sessions"][session_id] = session
    save_state(state)
    for item in task.get("launch", []):
        launch_item(item)
    launch_task_resources(task)
    if config["signals"]["screenshot_at_start"]:
        capture_screenshot(session_dir / "start.png")
    total_sec = max(1, int(duration_min * 60))
    interval = int(config["signals"]["window_sample_interval_sec"])
    midpoint_done = False
    start = time.time()
    while time.time() - start < total_sec:
        session["samples"].append(frontmost_window())
        elapsed = time.time() - start
        if config["signals"]["screenshot_midpoint"] and not midpoint_done and elapsed >= total_sec / 2:
            capture_screenshot(session_dir / "midpoint.png")
            midpoint_done = True
        time.sleep(min(interval, max(0.2, total_sec - elapsed)))
    if config["signals"]["screenshot_at_end"]:
        capture_screenshot(session_dir / "end.png")
    observed_elapsed_min = (time.time() - start) / 60
    (session_dir / "window_samples.json").write_text(
        json.dumps(session["samples"], ensure_ascii=False, indent=2),
        encoding="utf-8"
    )
    signal_summary = compute_signal_score(task, session["samples"], session_dir, observed_elapsed_min)
    (session_dir / "signal_summary.json").write_text(
        json.dumps(signal_summary, ensure_ascii=False, indent=2),
        encoding="utf-8"
    )
    threshold = config["completion_policy"]["candidate_threshold"]
    recommendation = "Candidate Complete" if signal_summary["score"] >= threshold else "Needs Review"
    session["state"] = "Needs Review"
    session["recommendation"] = recommendation
    session["ended_at"] = dt.datetime.now().isoformat(timespec="seconds")
    session["signal_summary"] = signal_summary
    set_task_state(state, args.task, recommendation)
    if recommendation == "Candidate Complete":
        set_task_state(state, args.task, "Needs Review")
    packet = build_review_packet(session, task, signal_summary, session_dir)
    session["review_packet"] = packet
    state["sessions"][session_id] = session
    save_state(state)
    print(json.dumps({"session_id": session_id, "review_packet": packet, "recommendation": recommendation}, ensure_ascii=False, indent=2))


def apply_review(session_id, status, notes="", adjust_min=None, review_source="manual_cli"):
    config = load_config()
    state = load_state()
    session = state["sessions"].get(session_id)
    if not session:
        raise SystemExit(f"Unknown session id: {session_id}")
    task_id = session["task_id"]
    task = state["tasks"][task_id]
    mapping = {
        "completed": "Completed",
        "partially_completed": "Partially Completed",
        "not_completed": "Failed",
        "failed": "Failed",
        "blocked": "Blocked"
    }
    final_state = mapping[status]
    set_task_state(state, task_id, final_state)
    session["uone_review"] = {
        "status": status,
        "notes": notes or "",
        "adjust_min": adjust_min,
        "source": review_source,
        "reviewed_at": dt.datetime.now().isoformat(timespec="seconds")
    }
    session["state"] = final_state
    key = task["task_key"]
    current_min = int(task["target_min"])
    adaptive = state.setdefault("adaptive_time", {})
    if final_state == "Completed":
        adaptive.pop(key, None)
    elif final_state in {"Partially Completed", "Failed"}:
        bump = int(adjust_min if adjust_min is not None else config["completion_policy"]["time_escalation_min"])
        cap = int(config["completion_policy"]["max_target_min"])
        adaptive[key] = min(cap, current_min + bump)
    state["sessions"][session_id] = session
    save_state(state)
    result = {"task_id": task_id, "final_state": final_state, "next_target_min": state.get("adaptive_time", {}).get(key)}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return result


def review_session(args):
    apply_review(args.session, args.status, args.notes or "", args.adjust_min, "manual_cli")


def parse_uone_review_text(text):
    normalized = re.sub(r"\s+", "", text.lower())
    adjust_min = None
    m = re.search(r"(?:多定|加|增加|延长|再来|再做|追加)(\d{1,3})(?:分钟|min|m)?", normalized)
    if not m:
        m = re.search(r"(\d{1,3})(?:分钟|min|m)", normalized)
    if m:
        adjust_min = int(m.group(1))
    if any(word in normalized for word in ["完成了", "已完成", "通过", "可以过", "算完成", "done", "completed"]):
        status = "completed"
    elif any(word in normalized for word in ["部分完成", "做了一部分", "差一点", "快完成", "partial"]):
        status = "partially_completed"
    elif any(word in normalized for word in ["卡住", "打不开", "无法", "没法", "blocked"]):
        status = "blocked"
    elif any(word in normalized for word in ["没有完成", "没完成", "未完成", "不通过", "重做", "notcompleted", "failed"]):
        status = "not_completed"
    elif adjust_min is not None:
        status = "not_completed"
    else:
        raise SystemExit("Could not infer uone status from text. Use clearer words: 完成了 / 没完成 / 部分完成 / 卡住, optionally 多定20分钟.")
    return status, adjust_min


def review_text_session(args):
    text = args.text or ""
    if args.file:
        text = Path(args.file).read_text(encoding="utf-8")
    status, adjust_min = parse_uone_review_text(text)
    apply_review(args.session, status, text, adjust_min, "uone_text_or_ocr")


def list_tasks(args):
    state = load_state()
    for task_id, task in state["tasks"].items():
        if args.date and task.get("date") != args.date:
            continue
        print(f"{task_id}\t{task.get('state')}\t{task.get('target_min')}m\t{task.get('title')}")


def material_report(args):
    state = load_state()
    date_text = args.date
    tasks = [t for t in state["tasks"].values() if not date_text or t.get("date") == date_text]
    for task in tasks:
        print(f"\n{task['id']}\n{task['title']}")
        for resource in task.get("resources", []):
            label = resource.get("label")
            target = resource.get("target")
            if args.core_only and not (
                str(label).startswith("resource_index_")
                or
                str(label).startswith("task_excerpt")
                or label in {"question_file", "supplemental_courseware"}
            ):
                continue
            exists = target_exists(target)
            page_range = resource.get("page_range")
            page_text = f" p{page_range[0]}-{page_range[1]}" if page_range else ""
            index_name = f" ({resource.get('index_filename')})" if resource.get("missing") else ""
            print(f"- {label}{page_text}{index_name}: {'OK' if exists else 'MISSING'}")
            print(f"  {target or resource.get('index_unit') or ''}")


def resource_index_report(args):
    config = load_config()
    courses = [args.course] if args.course else ["AP_Calculus_BC", "AP_CSA"]
    summary = {"found": 0, "missing": 0}
    for course in courses:
        print(f"\n{course} 资源索引")
        for entry in resource_index_entries(config, course):
            target = resolve_index_resource(config, entry)
            canvas_url = canvas_assignment_url(config, entry) if not target else None
            if target:
                summary["found"] += 1
                print(f"OK\t{entry['category']}\t{entry['unit']}\t{entry['filename']}")
                if args.verbose:
                    print(f"  {target}")
            elif canvas_url:
                summary["found"] += 1
                print(f"CANVAS\t{entry['category']}\t{entry['unit']}\t{entry['filename']}")
                if args.verbose:
                    meta = canvas_assignment_metadata(config, entry)
                    print(f"  {canvas_url}")
                    print(f"  expected title: {meta.get('expected_title', '')}")
                    print(f"  title verified: {meta.get('title_verified', False)}")
            else:
                summary["missing"] += 1
                print(f"MISSING\t{entry['category']}\t{entry['unit']}\t{entry['filename']}")
                print(f"  index path: {entry.get('path_hint')}")
    print(json.dumps(summary, ensure_ascii=False))


def resource_text(resource, max_pages=3):
    target = resource.get("target")
    if not target or not Path(target).exists() or Path(target).suffix.lower() != ".pdf":
        return ""
    try:
        from pypdf import PdfReader
        reader = PdfReader(target)
        chunks = []
        for page in reader.pages[:max_pages]:
            chunks.append(page.extract_text() or "")
        return " ".join(" ".join(chunks).split())
    except Exception:
        return ""


def audit_task_materials(task):
    issues = []
    warnings = []
    resources = task.get("resources", [])
    course = task.get("course")
    title = task.get("title", "")
    kind = task.get("kind", "")
    indexed = [r for r in resources if str(r.get("label", "")).startswith("resource_index_")]
    if course in {"AP_CSA", "AP_Calculus_BC"} and not indexed:
        issues.append("no resource-index material resolved for this task")
    for resource in resources:
        target = resource.get("target")
        if resource.get("missing"):
            issues.append(f"missing resource-index file: {resource.get('index_filename')}")
        if target and not target_exists(target):
            issues.append(f"missing file: {resource.get('label')}")
    if course == "AP_CSA" and kind == "CSA_CONCEPT":
        opened_labels = {
            "resource_index_syllabus",
            "resource_index_textbook",
            "resource_index_practice",
            "resource_index_checklist",
            "question_file",
        }
        if any(r.get("label") == "task_excerpt_java_illuminated" for r in resources):
            # Java reference is allowed, but it must not be part of launch_task_resources.
            pass
        if not any(r.get("label") in opened_labels for r in resources):
            issues.append("CSA task has no launchable primary material")
    return issues, warnings


def audit_materials(args):
    config = load_config()
    audit_state = {"tasks": {}, "sessions": {}, "adaptive_time": {}}
    start = parse_date_arg(args.start) if args.start else today_date()
    checked = 0
    failed = 0
    warned = 0
    for offset in range(args.days):
        current_date = start + dt.timedelta(days=offset)
        tasks = []
        tasks.extend(build_bc_tasks(config, audit_state, current_date))
        tasks.extend(build_csa_tasks(config, audit_state, current_date))
        if not tasks:
            continue
        for task in tasks:
            if args.course and task.get("course") != args.course:
                continue
            checked += 1
            issues, warnings = audit_task_materials(task)
            if issues:
                failed += 1
                print(f"FAIL\t{current_date}\t{task['id']}\t{task['title']}")
                for issue in issues:
                    print(f"  - {issue}")
                for resource in task.get("resources", []):
                    label = resource.get("label")
                    if str(label).startswith("task_excerpt") or str(label).startswith("supplemental_print_packet") or label == "local_unit_page":
                        page_range = resource.get("page_range")
                        page_text = f" p{page_range[0]}-{page_range[1]}" if page_range else ""
                        print(f"    {label}{page_text}: {resource.get('target')}")
            elif warnings:
                warned += 1
                print(f"WARN\t{current_date}\t{task['id']}\t{task['title']}")
                for warning in warnings:
                    print(f"  - {warning}")
            elif args.verbose:
                print(f"OK\t{current_date}\t{task['id']}\t{task['title']}")
    print(json.dumps({"checked": checked, "failed": failed, "warned": warned, "start": start.isoformat(), "days": args.days}, ensure_ascii=False))
    if failed:
        raise SystemExit(1)


def reset_task(args):
    state = load_state()
    if args.task not in state["tasks"]:
        raise SystemExit(f"Unknown task id: {args.task}")
    set_task_state(state, args.task, "Planned")
    state["tasks"][args.task]["last_session_id"] = None
    save_state(state)
    print(json.dumps({"task_id": args.task, "state": "Planned"}, ensure_ascii=False, indent=2))


def cmd_today(args):
    config = load_config()
    state = load_state()
    current_date = parse_date_arg(args.date) if args.date else today_date()
    tasks = generate_today(config, state, current_date)
    print(f"Wrote {len(tasks)} tasks to {BLACKBOARD_PATH}")
    for task in tasks:
        print(f"{task['id']}\t{task['target_min']}m\t{task['title']}")


def main():
    parser = argparse.ArgumentParser(description="AP Learning OS workflow engine")
    sub = parser.add_subparsers(dest="command", required=True)
    p_today = sub.add_parser("today", help="generate today's blackboard tasks")
    p_today.add_argument("--date", help="YYYY-MM-DD")
    p_today.set_defaults(func=cmd_today)
    p_start = sub.add_parser("start", help="start and monitor a task session")
    p_start.add_argument("--task", required=True)
    p_start.add_argument("--duration-min", type=float)
    p_start.set_defaults(func=start_session)
    p_review = sub.add_parser("review", help="apply uone final review")
    p_review.add_argument("--session", required=True)
    p_review.add_argument("--status", required=True, choices=["completed", "partially_completed", "not_completed", "failed", "blocked"])
    p_review.add_argument("--notes")
    p_review.add_argument("--adjust-min", type=int, help="minutes to add to the next attempt when not completed")
    p_review.set_defaults(func=review_session)
    p_review_text = sub.add_parser("review-text", help="apply uone review from chat/OCR text")
    p_review_text.add_argument("--session", required=True)
    p_review_text.add_argument("--text", help="uone chat/OCR text")
    p_review_text.add_argument("--file", help="file containing uone chat/OCR text")
    p_review_text.set_defaults(func=review_text_session)
    p_list = sub.add_parser("tasks", help="list known tasks")
    p_list.add_argument("--date")
    p_list.set_defaults(func=list_tasks)
    p_materials = sub.add_parser("materials", help="report task resource mappings")
    p_materials.add_argument("--date")
    p_materials.add_argument("--core-only", action="store_true")
    p_materials.set_defaults(func=material_report)
    p_resolve = sub.add_parser("resolve-resources", help="resolve workbook resource-index files on this Mac")
    p_resolve.add_argument("--course", choices=["AP_Calculus_BC", "AP_CSA"])
    p_resolve.add_argument("--verbose", action="store_true")
    p_resolve.set_defaults(func=resource_index_report)
    p_audit = sub.add_parser("audit-materials", help="audit future task material mappings without changing state")
    p_audit.add_argument("--start", help="YYYY-MM-DD")
    p_audit.add_argument("--days", type=int, default=60)
    p_audit.add_argument("--course", choices=["AP_CSA", "AP_Calculus_BC"])
    p_audit.add_argument("--verbose", action="store_true")
    p_audit.set_defaults(func=audit_materials)
    p_reset = sub.add_parser("reset-task", help="reset a task to Planned after accidental/test start")
    p_reset.add_argument("--task", required=True)
    p_reset.set_defaults(func=reset_task)
    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
